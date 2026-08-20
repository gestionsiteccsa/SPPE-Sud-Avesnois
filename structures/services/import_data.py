import csv
import io
import re
import unicodedata
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from django.contrib.auth.base_user import AbstractBaseUser
from django.core.exceptions import ValidationError
from django.db import transaction

from communes.models import Commune
from structures.audit import audit_actor
from structures.models import AuditLog, JOURS_SEM, Structure, TypeStructure


MAX_IMPORT_ROWS = 5_000
REQUIRED_COLUMNS = frozenset({"NOM"})


class ImportDataError(ValueError):
    """Erreur de validation destinée à être présentée sans détail technique."""


def _normalize_name(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    stripped = "".join(c for c in normalized if not unicodedata.combining(c))
    return " ".join(stripped.lower().split())


@dataclass(frozen=True, slots=True)
class PreparedStructure:
    values: dict[str, object]
    type_name: str
    commune_name: str
    postal_code: str


def _as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized_row(row: Mapping[str, object]) -> dict[str, str]:
    return {_as_text(key): _as_text(value) for key, value in row.items() if key is not None}


def _parse_bool(value: str, *, field_label: str) -> bool | None:
    normalized = value.strip().lower()
    if normalized in {"", "non concerné", "non concerne", "nc"}:
        return None
    if normalized in {"oui", "o", "yes", "ok", "vrai", "1"}:
        return True
    if normalized in {"non", "n", "no", "nnon", "faux", "0"}:
        return False
    raise ImportDataError(f"Valeur invalide pour « {field_label} » : {value!r}.")


def _parse_optional_integer(value: str, *, field_label: str) -> int | None:
    if value.strip() == "":
        return None
    try:
        parsed = int(value)
    except ValueError as error:
        raise ImportDataError(
            f"« {field_label} » doit être un nombre entier : {value!r}."
        ) from error
    if parsed < 0:
        raise ImportDataError(f"« {field_label} » ne peut pas être négatif.")
    return parsed


def _parse_places(value: str) -> tuple[int | None, bool, bool]:
    normalized = value.strip().lower()
    if normalized in {"", "non communiqué", "non communique", "nc"}:
        return None, False, True
    if normalized == "complet":
        return None, True, False
    return (
        _parse_optional_integer(value, field_label="places disponibles"),
        False,
        False,
    )


def _parse_date(value: str) -> date | None:
    if value.strip() == "":
        return None
    for date_format in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value.strip(), date_format).date()
        except ValueError:
            continue
    raise ImportDataError(f"Date invalide : {value!r}. Formats acceptés : JJ/MM/AAAA ou AAAA-MM-JJ.")


def _age_unit(value: str) -> str:
    if value.startswith("semaine"):
        return "semaine"
    if value.startswith("mois"):
        return "mois"
    return "ans"


def _parse_age(value: str) -> dict[str, object]:
    if value.strip() == "":
        return {}
    normalized = value.strip().lower()
    if normalized in {"non renseigné", "non renseigne", "nr"}:
        return {"age_non_renseigne": True}

    result: dict[str, object] = {"age_non_renseigne": False}
    parts = re.split(r"\s*[-–]\s*", normalized)
    age_pattern = re.compile(r"(\d+)\s*(semaine|semaines|mois|an|ans)?")
    if len(parts) == 2:
        minimum = age_pattern.fullmatch(parts[0].strip())
        maximum = age_pattern.fullmatch(parts[1].strip())
        if maximum is None or maximum.group(2) is None or minimum is None:
            raise ImportDataError(f"Tranche d'âge invalide : {value!r}.")
        result.update(
            age_min=int(minimum.group(1)),
            age_min_unite=_age_unit(minimum.group(2) or maximum.group(2)),
            age_max=int(maximum.group(1)),
            age_max_unite=_age_unit(maximum.group(2)),
        )
        return result

    match = age_pattern.fullmatch(normalized)
    if match is None or match.group(2) is None:
        raise ImportDataError(f"Tranche d'âge invalide : {value!r}.")
    result.update(age_max=int(match.group(1)), age_max_unite=_age_unit(match.group(2)))
    return result


def parser_horaires(raw: str) -> list[dict[str, object]]:
    """Convertit la plage simple utilisée par les fichiers source en JSON métier."""
    closed_days = [{"jour": day, "ferme": True} for day, _label in JOURS_SEM]
    if raw.strip() == "":
        return closed_days

    aliases = {
        "lun": "lundi",
        "mar": "mardi",
        "mer": "mercredi",
        "jeu": "jeudi",
        "ven": "vendredi",
        "sam": "samedi",
        "dim": "dimanche",
    }
    match = re.match(
        r"(\w+)\s*[-–/]\s*(\w+)\s+(\d{1,2})h(\d{2})?\s*[-–]\s*(\d{1,2})h(\d{2})?",
        raw.strip(),
        flags=re.IGNORECASE,
    )
    if match is None:
        return closed_days

    day_names = [day for day, _label in JOURS_SEM]
    first_day = aliases.get(match.group(1).lower()[:3])
    last_day = aliases.get(match.group(2).lower()[:3])
    if first_day not in day_names or last_day not in day_names:
        return closed_days

    first_index = day_names.index(first_day)
    last_index = day_names.index(last_day)
    if first_index > last_index:
        return closed_days

    opening = f"{int(match.group(3)):02d}:{match.group(4) or '00'}"
    closing = f"{int(match.group(5)):02d}:{match.group(6) or '00'}"
    schedule = []
    for index, day in enumerate(day_names):
        if first_index <= index <= last_index:
            schedule.append(
                {
                    "jour": day,
                    "ferme": False,
                    "ouverture": opening,
                    "fermeture": closing,
                }
            )
        else:
            schedule.append({"jour": day, "ferme": True})
    return schedule


def _prepare_row(row: Mapping[str, object], *, line_number: int) -> PreparedStructure | None:
    normalized = _normalized_row(row)
    if not any(normalized.values()):
        return None
    name = normalized.get("NOM", "").strip()
    if name == "":
        raise ImportDataError(f"Ligne {line_number} : la colonne « NOM » est obligatoire.")

    try:
        available_places, complete, not_communicated = _parse_places(
            normalized.get("places dispos", "")
        )
        raw_reference = normalized.get("Référencé monenfant.fr", "")
        referenced = _parse_bool(
            "non" if raw_reference.strip().lower() in {"non concerné", "non concerne"} else raw_reference,
            field_label="Référencé monenfant.fr",
        )
        raw_schedule = normalized.get("Horaires", "")
        values: dict[str, object] = {
            "nom_structure": name,
            "reference_monenfant": True if raw_reference.strip() == "" else bool(referenced),
            "date_mise_a_jour_monenfant": _parse_date(normalized.get("Date mise à jour", "")),
            "adresse": normalized.get("adresse", ""),
            "telephone": normalized.get("tel", ""),
            "email": normalized.get("email", ""),
            "places_disponibles": available_places,
            "places_complet": complete,
            "places_non_communique": not_communicated,
            "conditions_places": normalized.get("conditions places dispo", ""),
            "nb_places_total": _parse_optional_integer(
                normalized.get("nb de places total", ""),
                field_label="nombre de places total",
            ),
            "horaires": parser_horaires(raw_schedule),
            "horaires_notes": raw_schedule,
            "accueil_handicap": _parse_bool(
                normalized.get("Accueil handicap", ""),
                field_label="Accueil handicap",
            ),
            "site_web": "" if normalized.get("Site web", "").lower() == "monenfant.fr" else normalized.get("Site web", ""),
            "directeur": normalized.get("Directrice.eur", ""),
            "tel_direction": normalized.get("tel direction", ""),
            "email_direction": normalized.get("email direction", ""),
            "statut": normalized.get("statut", ""),
            "aides": normalized.get("Aides", ""),
            "nb_professionnels": _parse_optional_integer(
                normalized.get("Nb de professionnel.les", ""),
                field_label="nombre de professionnel·les",
            ),
            "recrutement": _parse_bool(
                normalized.get("recrutement en cours ?", ""),
                field_label="recrutement en cours",
            ),
            "accueil_urgence": _parse_bool(
                normalized.get("OK accueil d'urgence", ""),
                field_label="accueil d'urgence",
            ),
        }
        values.update(_parse_age(normalized.get("Tranche d'âge", "")))
        Structure(**values).full_clean(exclude={"type", "commune"})
    except (ImportDataError, ValidationError) as error:
        details = "; ".join(error.messages) if isinstance(error, ValidationError) else str(error)
        raise ImportDataError(f"Ligne {line_number} ({name}) : {details}") from error

    return PreparedStructure(
        values=values,
        type_name=normalized.get("Type", ""),
        commune_name=normalized.get("commune", ""),
        postal_code=normalized.get("code postal", ""),
    )


def _validate_headers(headers: Iterable[object]) -> None:
    normalized = {_as_text(header) for header in headers}
    missing = REQUIRED_COLUMNS - normalized
    if missing:
        raise ImportDataError(f"Colonne obligatoire absente : {', '.join(sorted(missing))}.")


def _collect_bounded_rows(rows: Iterable[Mapping[str, object]]) -> list[dict[str, str]]:
    collected = []
    for index, row in enumerate(rows, start=1):
        if index > MAX_IMPORT_ROWS:
            raise ImportDataError(f"Le fichier dépasse la limite de {MAX_IMPORT_ROWS} lignes.")
        collected.append(
            {
                _as_text(key): _as_text(value)
                for key, value in row.items()
                if key is not None
            }
        )
    return collected


def read_csv_upload(upload) -> list[dict[str, str]]:
    try:
        decoded = upload.read().decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ImportDataError("Le fichier CSV doit être encodé en UTF-8.") from error
    reader = csv.DictReader(io.StringIO(decoded))
    if reader.fieldnames is None:
        raise ImportDataError("Le fichier CSV ne contient pas d'en-tête.")
    _validate_headers(reader.fieldnames)
    return _collect_bounded_rows(reader)


def read_xlsx_upload(upload) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
    except Exception as error:
        raise ImportDataError("Le fichier XLSX est illisible ou invalide.") from error
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            raise ImportDataError("Le fichier XLSX est vide.")
        _validate_headers(header)
        headers = [_as_text(cell) for cell in header]
        rows = (
            {
                column: _as_text(row[index]) if index < len(row) else ""
                for index, column in enumerate(headers)
            }
            for row in iterator
        )
        return _collect_bounded_rows(rows)
    finally:
        workbook.close()


def read_rows_from_path(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise ImportDataError(f"Fichier introuvable : {path}.")
    if path.suffix.lower() == ".csv":
        with path.open("rb") as csv_file:
            return read_csv_upload(csv_file)
    if path.suffix.lower() == ".xlsx":
        with path.open("rb") as xlsx_file:
            return read_xlsx_upload(xlsx_file)
    raise ImportDataError("Format non supporté. Utilisez un fichier .csv ou .xlsx.")


def import_rows(
    rows: Iterable[Mapping[str, object]],
    *,
    replace: bool,
    actor: AbstractBaseUser | None,
) -> int:
    row_list = list(rows)
    if len(row_list) > MAX_IMPORT_ROWS:
        raise ImportDataError(f"Le fichier dépasse la limite de {MAX_IMPORT_ROWS} lignes.")

    prepared = []
    for line_number, row in enumerate(row_list, start=2):
        item = _prepare_row(row, line_number=line_number)
        if item is not None:
            prepared.append(item)
    if not prepared:
        raise ImportDataError("Le fichier ne contient aucune structure exploitable.")

    skipped = 0
    with transaction.atomic(), audit_actor(actor):
        if replace:
            Structure.objects.all().delete()
        for item in prepared:
            type_object = None
            if item.type_name:
                type_object, _created = TypeStructure.objects.get_or_create(nom=item.type_name)

            commune = None
            if item.commune_name:
                commune, _created = Commune.objects.get_or_create(
                    nom=item.commune_name,
                    defaults={"code_postal": item.postal_code},
                )
            elif item.postal_code:
                commune = Commune.objects.filter(code_postal=item.postal_code).first()

            if not replace and commune is not None:
                identity = item.values.get("nom_structure", "").strip()
                if identity:
                    duplicates = Structure.objects.filter(commune=commune)
                    candidate_normalized = _normalize_name(identity)
                    already_exists = any(
                        _normalize_name(existing.nom_structure) == candidate_normalized
                        for existing in duplicates
                    )
                    if already_exists:
                        skipped += 1
                        continue

            structure = Structure(type=type_object, commune=commune, **item.values)
            structure.full_clean()
            structure.save()

        if actor is not None:
            AuditLog.objects.create(
                user=actor,
                action="import",
                model_name="Structure",
                object_repr=f"Import de {len(prepared)} structure(s)",
                changes={
                    "count": len(prepared),
                    "replace": replace,
                    "ignored_duplicates": skipped,
                },
            )
    return len(prepared) - skipped
